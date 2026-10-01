from io import BytesIO
from typing import Any

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.api.admin_router import require_admin
from app.database import get_db
from app.domain import models
from app.services.inventory_excel_service import (
    commit_workbook,
    validate_workbook,
)


router = APIRouter(
    prefix="/admin-panel/inventory-excel",
    tags=["Admin - Inventory Excel Import"],
)

MAX_FILE_SIZE = 10 * 1024 * 1024


async def _read_uploaded_excel(file: UploadFile) -> bytes:
    filename = (file.filename or "").strip()

    if not filename.lower().endswith(".xlsx"):
        raise HTTPException(
            status_code=400,
            detail="فقط فایل Excel با پسوند xlsx مجاز است.",
        )

    content = await file.read()

    if not content:
        raise HTTPException(
            status_code=400,
            detail="فایل انتخاب‌شده خالی است.",
        )

    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=413,
            detail="حجم فایل نباید بیشتر از ۱۰ مگابایت باشد.",
        )

    return content


def _autosize_sheet(sheet) -> None:
    for column_cells in sheet.columns:
        maximum_length = 0
        column_letter = get_column_letter(column_cells[0].column)

        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            maximum_length = max(maximum_length, len(value))

        sheet.column_dimensions[column_letter].width = min(
            maximum_length + 3,
            35,
        )


@router.get("/template")
def download_inventory_excel_template(
    current_user: models.User = Depends(require_admin),
):
    workbook = Workbook()

    inventory_sheet = workbook.active
    inventory_sheet.title = "Inventory"
    inventory_sheet.sheet_view.rightToLeft = True

    inventory_headers = [
        "نام انبار",
        "نام کالا",
        "کد کالا",
        "واحد پایه",
        "نوع کالا",
        "شماره بچ",
        "تاریخ تولید",
        "تاریخ انقضا",
        "موجودی کل",
        "موجودی رزرو",
        "ترکیبی",
        "وزن واحد گرم",
        "کد بسته‌بندی اصلی",
    ]

    inventory_sheet.append(inventory_headers)

    packaging_sheet = workbook.create_sheet("PackagingInventory")
    packaging_sheet.sheet_view.rightToLeft = True

    packaging_headers = [
        "نام انبار",
        "نام قلم بسته‌بندی",
        "کد بسته‌بندی",
        "نوع بسته‌بندی",
        "گنجایش گرم",
        "توضیحات",
        "تعداد کل",
        "تعداد رزرو",
    ]

    packaging_sheet.append(packaging_headers)

    guide_sheet = workbook.create_sheet("راهنما")
    guide_sheet.sheet_view.rightToLeft = True
    guide_sheet.append(["موضوع", "توضیح"])
    guide_sheet.append([
        "نام انبار",
        "باید دقیقاً با نام انبار فعال در سامانه یکسان باشد.",
    ])
    guide_sheet.append([
        "نوع کالا",
        "فقط raw_material یا finished_good",
    ])
    guide_sheet.append([
        "نوع بسته‌بندی",
        "container, lid, label, inner_box, carton, pallet, other",
    ])
    guide_sheet.append([
        "کد بسته‌بندی اصلی",
        (
            "برای هر کالا الزامی است؛ باید به یک قلم container "
            "با گنجایش حداقل برابر وزن واحد کالا اشاره کند."
        ),
    ])
    guide_sheet.append([
        "قانون مصرف بسته‌بندی",
        (
            "به ازای انتقال هر واحد کالا، یک واحد بسته‌بندی "
            "در مرحله dispatch از انبار مبدأ کسر می‌شود."
        ),
    ])
    guide_sheet.append([
        "تاریخ‌ها",
        "قالب پیشنهادی: YYYY-MM-DD",
    ])
    guide_sheet.append([
        "ترکیبی",
        "بله/خیر یا true/false",
    ])
    guide_sheet.append([
        "رفتار فایل تکراری",
        "موجودی همان انبار، کالا و بچ به‌روزرسانی می‌شود.",
    ])

    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)

        sheet.freeze_panes = "A2"
        _autosize_sheet(sheet)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    headers = {
        "Content-Disposition": (
            'attachment; filename="initial_inventory_template.xlsx"'
        )
    }

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers=headers,
    )


@router.post("/preview")
async def preview_inventory_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(require_admin),
) -> dict[str, Any]:
    content = await _read_uploaded_excel(file)

    try:
        result = validate_workbook(db, content)
    except ValueError as exc:
        raise HTTPException(
            status_code=400,
            detail=str(exc),
        )

    return {
        "valid": result["valid"],
        "summary": result["summary"],
        "errors": result["errors"],
        "preview_rows": result["preview_rows"],
    }


@router.post("/commit")
async def commit_inventory_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(require_admin),
) -> dict[str, Any]:
    content = await _read_uploaded_excel(file)

    try:
        result = commit_workbook(
            db=db,
            current_user=current_user,
            file_bytes=content,
        )
    except ValueError as exc:
        raise HTTPException(
            status_code=400,
            detail=str(exc),
        )
    except Exception:
        raise HTTPException(
            status_code=500,
            detail="ثبت اطلاعات Excel انجام نشد و تمام تغییرات برگشت داده شد.",
        )

    if not result["valid"]:
        raise HTTPException(
            status_code=400,
            detail={
                "message": "فایل دارای خطا است و هیچ اطلاعاتی ثبت نشد.",
                "summary": result["summary"],
                "errors": result["errors"],
            },
        )

    return result